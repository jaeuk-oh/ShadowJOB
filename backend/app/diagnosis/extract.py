"""이력서/포폴 파일 → 텍스트 추출 (P1.5+).

서버 단일 소스로 추출 → 붙여넣기와 동일한 진단 경로 재사용.
지원: .pdf(pypdf) · .docx(python-docx) · .txt/.md(decode).
.doc(구형 바이너리)는 미지원 → UnsupportedFileType.
"""
from __future__ import annotations

import io
from pathlib import Path

MAX_BYTES = 10 * 1024 * 1024  # 10MB
SUPPORTED = (".pdf", ".docx", ".txt", ".md")


class UnsupportedFileType(Exception):
    pass


class EmptyExtraction(Exception):
    pass


class FileTooLarge(Exception):
    pass


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    # 표 안의 텍스트도 수집
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts)


def _extract_txt(data: bytes) -> str:
    for enc in ("utf-8", "utf-16", "cp949", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_text(filename: str, data: bytes) -> str:
    """파일명 확장자로 디스패치해 텍스트 추출. 실패 시 명확한 예외."""
    if len(data) > MAX_BYTES:
        raise FileTooLarge(f"파일이 너무 큽니다(>{MAX_BYTES // (1024 * 1024)}MB)")

    ext = Path(filename).suffix.lower()
    if ext == ".doc":
        raise UnsupportedFileType(
            ".doc(구형 Word)는 지원하지 않습니다. .docx 또는 .pdf로 저장해 올려주세요."
        )
    if ext not in SUPPORTED:
        raise UnsupportedFileType(
            f"지원하지 않는 형식입니다: {ext or '(확장자 없음)'} — {', '.join(SUPPORTED)}만 가능"
        )

    if ext == ".pdf":
        text = _extract_pdf(data)
    elif ext == ".docx":
        text = _extract_docx(data)
    else:
        text = _extract_txt(data)

    text = text.strip()
    if not text:
        raise EmptyExtraction(
            "파일에서 텍스트를 추출하지 못했습니다(이미지 기반 PDF일 수 있음). 텍스트로 붙여넣어 주세요."
        )
    return text
